from __future__ import annotations

from datetime import datetime
from pathlib import Path
from PySide6.QtCore import QUrl, Signal, Qt
from PySide6.QtGui import QDesktopServices
from PySide6.QtWidgets import QBoxLayout, QCheckBox, QComboBox, QDialog, QDialogButtonBox, QFileDialog, QFormLayout, QFrame, QHBoxLayout, QLabel, QLineEdit, QMessageBox, QPushButton, QVBoxLayout, QWidget

from ..constants import APP_VERSION, MONTHS
from ..importers import import_legacy_finance_workbook, import_mercadopago_report
from ..secure_store import delete_token, get_token, save_token
from ..utils import money
from ..layouts import responsive_mode
from .common import page_header, scroll_container

class LegacyImportDialog(QDialog):
    def __init__(self, db, payload, parent=None):
        super().__init__(parent)
        self.db = db
        self.payload = payload
        self.setWindowTitle("Importar historial del Excel")
        self.setMinimumWidth(560)
        root = QVBoxLayout(self); root.setContentsMargins(22,20,22,20); root.setSpacing(13)

        title = QLabel("Importar tu historial anterior"); title.setObjectName("DialogTitle")
        desc = QLabel(
            "Voy a traer las categorías y cada importe mensual de Gastos e Ingresos como datos visibles dentro de App Gastos. "
            "Los vas a encontrar también en Movimientos. Como el Excel no guardaba el día ni la cuenta, esos dos datos no se inventan."
        ); desc.setObjectName("Muted"); desc.setWordWrap(True)
        root.addWidget(title); root.addWidget(desc)

        months = payload.get("months", [])
        first = f"{MONTHS[months[0][1]-1]} {months[0][0]}" if months else "—"
        last = f"{MONTHS[months[-1][1]-1]} {months[-1][0]}" if months else "—"
        card = QFrame(); card.setObjectName("Card")
        cl = QVBoxLayout(card); cl.setContentsMargins(16,14,16,14); cl.setSpacing(5)
        lines = [
            f"Archivo: {payload.get('file_name','')}",
            f"Período detectado: {first} → {last}",
            f"Categorías/subcategorías con datos: {payload.get('category_count',0)}",
            f"Gastos detectados: {money(payload.get('expense_total',0), db.currency_symbol(), db.balances_hidden())}",
            f"Ingresos detectados: {money(payload.get('income_total',0), db.currency_symbol(), db.balances_hidden())}",
        ]
        for text in lines:
            lab=QLabel(text); lab.setObjectName("Muted"); cl.addWidget(lab)
        root.addWidget(card)

        self.closed = QCheckBox("Importar solo meses cerrados (recomendado)")
        self.closed.setChecked(True)
        self.closed.setToolTip("Evita traer como gasto real cuotas o importes que el Excel ya tenía cargados para el mes actual o meses futuros.")
        self.skip_real = QCheckBox("Omitir meses que ya tengan movimientos reales en App Gastos")
        self.skip_real.setChecked(True)
        self.skip_history = QCheckBox("Omitir meses que ya tengan historial importado")
        self.skip_history.setChecked(True)
        root.addWidget(self.closed); root.addWidget(self.skip_real); root.addWidget(self.skip_history)

        note=QLabel("Antes de importar se crea un backup automático de la base.")
        note.setObjectName("SmallMuted"); root.addWidget(note)

        buttons=QDialogButtonBox(QDialogButtonBox.StandardButton.Cancel)
        accept=QPushButton("Importar datos del Excel"); accept.clicked.connect(self.accept)
        buttons.addButton(accept,QDialogButtonBox.ButtonRole.AcceptRole)
        buttons.rejected.connect(self.reject); root.addWidget(buttons)

    def options(self):
        return {
            "skip_transaction_months": self.skip_real.isChecked(),
            "skip_existing_history_months": self.skip_history.isChecked(),
            "closed_months_only": self.closed.isChecked(),
        }


class SettingsPage(QWidget):
    data_changed=Signal()
    theme_changed=Signal(str)
    ui_scale_changed=Signal(float)
    def __init__(self,db):
        super().__init__(); self.db=db; self._compact=False; outer=QVBoxLayout(self); outer.setContentsMargins(0,0,0,0); scroll,_,root=scroll_container(); outer.addWidget(scroll)
        root.addWidget(page_header("Ajustes","Preferencias, backups y portabilidad de tus datos"))
        pref=QFrame(); pref.setObjectName("Card"); pl=QVBoxLayout(pref); pl.setContentsMargins(22,20,22,20); pl.setSpacing(12)
        pref_head=QVBoxLayout(); pref_head.setSpacing(2)
        t=QLabel("Apariencia y accesibilidad"); t.setObjectName("SectionTitle")
        pref_sub=QLabel("Elegí el tema y ajustá el tamaño de toda la interfaz a tu monitor o comodidad visual."); pref_sub.setObjectName("SmallMuted"); pref_sub.setWordWrap(True)
        pref_head.addWidget(t); pref_head.addWidget(pref_sub); pl.addLayout(pref_head)
        form=QFormLayout(); form.setSpacing(12)
        self.theme=QComboBox()
        for label,value in (("Oscuro · Menta","dark_mint"),("Oscuro · Azul","dark_ocean"),("Oscuro · Violeta","dark_violet"),("Oscuro · Arena","dark_sunset"),("Claro · Menta","light_mint"),("Claro · Azul","light_ocean"),("Claro · Violeta","light_violet"),("Claro · Arena","light_sunset")):
            self.theme.addItem(label,value)
        current_theme = db.get_setting("theme","dark_mint")
        if current_theme == "dark": current_theme = "dark_mint"
        elif current_theme == "light": current_theme = "light_mint"
        i=self.theme.findData(current_theme); self.theme.setCurrentIndex(max(i,0))
        self._loaded_theme=current_theme
        self.icon_style = QComboBox()
        self.icon_style.addItem("Ilustrados · pastel", "illustrated")
        self.icon_style.addItem("Contorno · sobrios", "outline")
        self.icon_style.setCurrentIndex(max(0, self.icon_style.findData(db.get_setting("icon_style", "illustrated"))))
        form.addRow("Estilo de iconos", self.icon_style)
        self.ui_scale=QComboBox()
        for label,value in (("80% · Compacto",0.80),("90%",0.90),("100% · Normal",1.00),("110%",1.10),("125% · Cómodo",1.25),("140% · Grande",1.40),("160% · Muy grande",1.60),("180%",1.80),("200% · Máximo",2.00)):
            self.ui_scale.addItem(label,value)
        try: saved_scale=float(db.get_setting("ui_scale","1.00") or 1.0)
        except Exception: saved_scale=1.0
        scale_index=min(range(self.ui_scale.count()), key=lambda idx: abs(float(self.ui_scale.itemData(idx))-saved_scale))
        self.ui_scale.setCurrentIndex(scale_index)
        self._loaded_ui_scale=float(self.ui_scale.currentData())
        self.currency=QLineEdit(db.currency_symbol()); self.currency.setMaxLength(4)
        self.hide=QCheckBox("Ocultar importes en toda la aplicación"); self.hide.setChecked(db.balances_hidden())
        form.addRow("Tema / color",self.theme); form.addRow("Escala de interfaz",self.ui_scale); form.addRow("Símbolo de moneda",self.currency); form.addRow("Privacidad",self.hide)
        pl.addLayout(form)
        scale_hint=QLabel("La escala modifica texto, íconos, botones y espacios. Desde v0.34 se aplica en el momento, sin cerrar la app."); scale_hint.setObjectName("SmallMuted"); scale_hint.setWordWrap(True); pl.addWidget(scale_hint)
        save=QPushButton("Guardar preferencias"); save.clicked.connect(self.save_preferences); pl.addWidget(save,0,Qt.AlignmentFlag.AlignLeft); root.addWidget(pref)
        integ=QFrame(); integ.setObjectName("Card"); il=QVBoxLayout(integ); il.setContentsMargins(20,18,20,18); il.setSpacing(10)
        it=QLabel("Integraciones"); it.setObjectName("SectionTitle"); il.addWidget(it)
        mp_head=QHBoxLayout()
        mp_title=QLabel("Mercado Pago"); mp_title.setObjectName("IntegrationTitle"); mp_head.addWidget(mp_title)
        mp_head.addStretch()
        self.mp_status=QLabel(); self.mp_status.setObjectName("IntegrationStatus")
        mp_head.addWidget(self.mp_status); il.addLayout(mp_head)
        mp_desc=QLabel("Conectá tu cuenta para traer movimientos con el reporte oficial de 'Todas las transacciones'. Los movimientos nuevos llegan primero a Por revisar; el Access Token queda guardado en el almacén seguro de Windows y nunca dentro de SQLite."); mp_desc.setObjectName("Muted"); mp_desc.setWordWrap(True); il.addWidget(mp_desc)
        mp_form=QFormLayout(); mp_form.setSpacing(10)
        self.mp_account=QComboBox(); self._load_mp_accounts()
        self.mp_token=QLineEdit(); self.mp_token.setEchoMode(QLineEdit.EchoMode.Password); self.mp_token.setPlaceholderText("Token guardado en Windows" if get_token() else "APP_USR-…")
        self.mp_auto=QCheckBox("Sincronizar al abrir App Gastos")
        self.mp_auto.setChecked(db.get_setting("mercadopago_auto_sync","0")=="1")
        self.mp_days=QComboBox()
        for label,value in (("Últimos 7 días",7),("Últimos 15 días",15),("Últimos 30 días",30),("Últimos 60 días",60)):
            self.mp_days.addItem(label,value)
        try: saved_days=int(db.get_setting("mercadopago_sync_days","30") or 30)
        except Exception: saved_days=30
        di=self.mp_days.findData(saved_days); self.mp_days.setCurrentIndex(di if di>=0 else self.mp_days.findData(30))
        mp_form.addRow("Cuenta local",self.mp_account); mp_form.addRow("Access Token",self.mp_token)
        mp_form.addRow("Rango de sincronización",self.mp_days); mp_form.addRow("Automatización",self.mp_auto); il.addLayout(mp_form)
        self.mp_account.currentIndexChanged.connect(self._save_mp_preferences)
        self.mp_days.currentIndexChanged.connect(self._save_mp_preferences)
        self.mp_auto.stateChanged.connect(self._save_mp_preferences)
        self.mp_actions=QBoxLayout(QBoxLayout.Direction.LeftToRight)
        mp_actions=self.mp_actions
        save_mp=QPushButton("Guardar token"); save_mp.clicked.connect(self.save_mp_token)
        test_mp=QPushButton("Conectar y configurar"); test_mp.setObjectName("SecondaryButton"); test_mp.clicked.connect(self.test_mp_api)
        import_mp=QPushButton("Importar reporte"); import_mp.setObjectName("SecondaryButton"); import_mp.clicked.connect(self.import_mp_report)
        remove_mp=QPushButton("Quitar token"); remove_mp.setObjectName("GhostButton"); remove_mp.clicked.connect(self.remove_mp_token)
        for b in (save_mp,test_mp,import_mp,remove_mp): mp_actions.addWidget(b)
        mp_actions.addStretch(); il.addLayout(mp_actions)
        mp_hint=QLabel("La primera conexión crea automáticamente la configuración de reportes que Mercado Pago exige antes de poder generar archivos. Si la API no estuviera habilitada para tu cuenta, podés seguir usando Importar reporte. App Gastos evita duplicados usando el ID externo de cada operación."); mp_hint.setObjectName("SmallMuted"); mp_hint.setWordWrap(True); il.addWidget(mp_hint)
        self._refresh_mp_status()
        root.addWidget(integ)

        data=QFrame(); data.setObjectName("Card"); dl=QVBoxLayout(data); dl.setContentsMargins(20,18,20,18); dl.setSpacing(10)
        t=QLabel("Tus datos"); t.setObjectName("SectionTitle"); dl.addWidget(t)
        path=QLabel(str(db.path)); path.setObjectName("Muted"); path.setWordWrap(True); dl.addWidget(path)
        legacy_hint=QLabel("¿Venías usando tu Excel de gastos? Podés traer los totales mensuales sin convertirlos en movimientos con fechas inventadas."); legacy_hint.setObjectName("SmallMuted"); legacy_hint.setWordWrap(True); dl.addWidget(legacy_hint)
        self.data_actions=QBoxLayout(QBoxLayout.Direction.LeftToRight)
        row=self.data_actions
        legacy=QPushButton("Importar Excel anterior"); legacy.clicked.connect(self.import_legacy_excel)
        backup=QPushButton("Crear backup"); backup.setObjectName("SecondaryButton"); backup.clicked.connect(self.backup)
        restore=QPushButton("Restaurar backup"); restore.setObjectName("SecondaryButton"); restore.clicked.connect(self.restore)
        export=QPushButton("Exportar a Excel"); export.setObjectName("SecondaryButton"); export.clicked.connect(self.export_excel)
        folder=QPushButton("Abrir carpeta de datos"); folder.setObjectName("GhostButton"); folder.clicked.connect(lambda:QDesktopServices.openUrl(QUrl.fromLocalFile(str(db.path.parent))))
        [row.addWidget(x) for x in (legacy,backup,restore,export,folder)]; row.addStretch(); dl.addLayout(row); root.addWidget(data)
        info=QFrame(); info.setObjectName("Card"); il=QVBoxLayout(info); il.setContentsMargins(20,18,20,18); t=QLabel("Acerca de"); t.setObjectName("SectionTitle"); il.addWidget(t); x=QLabel(f"App Gastos v{APP_VERSION}\nPython + PySide6 + SQLite\nLa base se guarda fuera de la carpeta del programa para que futuras versiones no pisen tus datos."); x.setObjectName("Muted"); il.addWidget(x); root.addWidget(info); root.addStretch()
    def _apply_responsive(self):
        compact = responsive_mode(self.width(), self.height()) != "wide"
        if compact == self._compact:
            return
        self._compact = compact
        direction = QBoxLayout.Direction.TopToBottom if compact else QBoxLayout.Direction.LeftToRight
        self.mp_actions.setDirection(direction)
        self.data_actions.setDirection(direction)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self._apply_responsive()

    def save_preferences(self):
        from PySide6.QtWidgets import QApplication
        icon_style = self.icon_style.currentData()
        self.db.set_setting("icon_style", icon_style)
        app = QApplication.instance()
        app.setProperty("icon_style", icon_style)
        for widget in app.allWidgets():
            widget.update()
        theme=self.theme.currentData()
        scale=float(self.ui_scale.currentData() or 1.0)
        theme_changed = theme != getattr(self, "_loaded_theme", theme)
        scale_changed=abs(scale-float(getattr(self,"_loaded_ui_scale",1.0)))>0.001
        self.db.set_setting("theme",theme)
        self.db.set_setting("ui_scale",f"{scale:.2f}")
        self.db.set_setting("currency_symbol",self.currency.text().strip() or "$")
        self.db.set_setting("hide_balances","1" if self.hide.isChecked() else "0")
        self._loaded_ui_scale=scale
        self._loaded_theme=theme
        if theme_changed:
            self.theme_changed.emit(theme)
        if scale_changed:
            self.ui_scale_changed.emit(scale)
        self.data_changed.emit()
        message = f"Preferencias guardadas. Escala aplicada: {round(scale*100)}%." if scale_changed else "Preferencias guardadas."
        QMessageBox.information(self,"Ajustes",message)

    def _load_mp_accounts(self):
        self.mp_account.clear()
        saved=self.db.get_setting("mercadopago_account_id","")
        chosen=0
        for i,a in enumerate(self.db.accounts()):
            self.mp_account.addItem(a["name"],a["id"])
            if saved and str(a["id"])==str(saved): chosen=i
            elif not saved and "mercado pago" in a["name"].lower(): chosen=i
        if self.mp_account.count(): self.mp_account.setCurrentIndex(chosen)

    def _save_mp_preferences(self, *args):
        if hasattr(self,"mp_account") and self.mp_account.currentData():
            self.db.set_setting("mercadopago_account_id",self.mp_account.currentData())
        if hasattr(self,"mp_days") and self.mp_days.currentData():
            self.db.set_setting("mercadopago_sync_days",self.mp_days.currentData())
        if hasattr(self,"mp_auto"):
            self.db.set_setting("mercadopago_auto_sync","1" if self.mp_auto.isChecked() else "0")

    def _refresh_mp_status(self, configured: bool | None = None, message: str = ""):
        if not hasattr(self, "mp_status"):
            return
        has_token = bool(get_token())
        if configured is None and has_token and self.db.get_setting("mercadopago_reports_configured", "0") == "1":
            configured = True
        if configured is True:
            self.mp_status.setText("● Conectado")
            self.mp_status.setProperty("state", "ok")
            self.mp_status.setToolTip(message or "Mercado Pago conectado")
        elif configured is False and has_token:
            self.mp_status.setText("● Falta configurar")
            self.mp_status.setProperty("state", "warn")
            self.mp_status.setToolTip(message or "Token guardado; todavía falta configurar reportes")
        elif has_token:
            self.mp_status.setText("● Token guardado")
            self.mp_status.setProperty("state", "neutral")
            self.mp_status.setToolTip("Usá Conectar y configurar para verificar la API")
        else:
            self.mp_status.setText("● No conectado")
            self.mp_status.setProperty("state", "neutral")
            self.mp_status.setToolTip("Pegá un Access Token de producción")
        self.mp_status.style().unpolish(self.mp_status); self.mp_status.style().polish(self.mp_status)

    def save_mp_token(self):
        token=self.mp_token.text().strip()
        if not token:
            QMessageBox.warning(self,"Mercado Pago","Pegá un Access Token antes de guardarlo."); return
        try:
            save_token(token)
            self._save_mp_preferences()
            self.mp_token.clear(); self.mp_token.setPlaceholderText("Token guardado en Windows")
            self._refresh_mp_status()
            QMessageBox.information(self,"Mercado Pago","Token guardado. Ahora podés usar ‘Conectar y configurar’ para preparar los reportes.")
        except Exception as e: QMessageBox.critical(self,"Mercado Pago",str(e))

    def remove_mp_token(self):
        delete_token(); self.db.set_setting("mercadopago_reports_configured", "0"); self.mp_token.clear(); self.mp_token.setPlaceholderText("APP_USR-…")
        self._refresh_mp_status()
        QMessageBox.information(self,"Mercado Pago","Se quitó el token guardado.")

    def test_mp_api(self):
        token=self.mp_token.text().strip() or None
        if token:
            try:
                save_token(token)
                self.mp_token.clear(); self.mp_token.setPlaceholderText("Token guardado en Windows")
            except Exception as e:
                QMessageBox.critical(self,"Mercado Pago",str(e)); return
        from ..mercadopago import ensure_report_configuration
        result=ensure_report_configuration()
        self.db.set_setting("mercadopago_reports_configured", "1" if result.ok and result.configured else "0")
        self._refresh_mp_status(result.configured if result.ok else False, result.message)
        if result.ok and result.configured:
            QMessageBox.information(
                self,"Mercado Pago conectado",
                result.message + "\n\nYa podés ir a ‘Por revisar’ y usar ‘Sincronizar Mercado Pago’."
            )
        else:
            QMessageBox.warning(self,"Mercado Pago",result.message)

    def import_mp_report(self):
        if not self.mp_account.currentData(): QMessageBox.warning(self,"Mercado Pago","Seleccioná la cuenta local de Mercado Pago."); return
        path,_=QFileDialog.getOpenFileName(self,"Importar reporte de Mercado Pago","","Reportes (*.csv *.xlsx *.xlsm);;CSV (*.csv);;Excel (*.xlsx *.xlsm)")
        if not path:return
        try:
            rows=import_mercadopago_report(path)
            result=self.db.stage_imported_movements(rows,"mercado_pago",int(self.mp_account.currentData()))
            self.db.set_setting("mercadopago_account_id",self.mp_account.currentData()); self.db.set_setting("mercadopago_last_import",path)
            self.data_changed.emit()
            QMessageBox.information(self,"Mercado Pago",f"Nuevos: {result['inserted']}\nYa existentes: {result['duplicates']}\n\nRevisalos desde la sección ‘Por revisar’.")
        except Exception as e: QMessageBox.critical(self,"Mercado Pago",str(e))

    def refresh(self):
        if not hasattr(self,"mp_account"):
            return
        self._refresh_mp_status()
        current=self.mp_account.currentData()
        self._load_mp_accounts()
        if current:
            i=self.mp_account.findData(current)
            if i>=0:self.mp_account.setCurrentIndex(i)

    def import_legacy_excel(self):
        path,_=QFileDialog.getOpenFileName(
            self,"Importar Excel anterior",str(Path.home()),"Excel (*.xlsx *.xlsm)"
        )
        if not path:
            return
        try:
            payload=import_legacy_finance_workbook(path)
            if self.db.historical_import_exists(payload["file_hash"]):
                synced=self.db.sync_legacy_categories(payload)
                self.data_changed.emit()
                QMessageBox.information(
                    self,"Excel ya importado",
                    "Este mismo archivo ya estaba importado, así que no dupliqué importes.\n\n"
                    f"Sincronicé la estructura de categorías ({synced.get('created',0)} nuevas). "
                    "Desde esta versión los importes anteriores aparecen también dentro de Movimientos como registros mensuales editables."
                )
                return
            dialog=LegacyImportDialog(self.db,payload,self)
            if not dialog.exec():
                return
            backup_path=self.db.path.parent / f"backup_pre_excel_{datetime.now():%Y-%m-%d_%H%M%S}.db"
            self.db.create_backup(backup_path)
            result=self.db.import_monthly_history(payload,**dialog.options())
            self.db.set_setting("legacy_excel_last_import",path)
            self.data_changed.emit()

            def fmt_months(items):
                return ", ".join(f"{MONTHS[m-1][:3]} {y}" for y,m in items) if items else "ninguno"
            msg=(
                f"Historial importado correctamente.\n\n"
                f"Meses importados: {fmt_months(result['months'])}\n"
                f"Gastos: {money(result['expense_total'],self.db.currency_symbol(),self.db.balances_hidden())}\n"
                f"Ingresos: {money(result['income_total'],self.db.currency_symbol(),self.db.balances_hidden())}\n"
                f"Categorías nuevas creadas: {result.get('categories_created',0)}\n\n"
                f"Meses omitidos por tener movimientos reales: {fmt_months(result['skipped_transaction_months'])}\n"
                f"Meses actuales/futuros omitidos: {fmt_months(result['skipped_open_months'])}\n\n"
                "Ya podés recorrer esos meses desde Dashboard, Movimientos, Categorías y Análisis."
            )
            QMessageBox.information(self,"Historial importado",msg)
        except Exception as e:
            QMessageBox.critical(self,"No se pudo importar",str(e))

    def backup(self):
        default=f"AppGastos_backup_{datetime.now():%Y-%m-%d_%H%M}.db"; path,_=QFileDialog.getSaveFileName(self,"Guardar backup",str(Path.home()/default),"Base SQLite (*.db)")
        if path:
            try:self.db.create_backup(path); QMessageBox.information(self,"Backup","Backup creado correctamente.")
            except Exception as e: QMessageBox.critical(self,"Backup",str(e))
    def restore(self):
        path,_=QFileDialog.getOpenFileName(self,"Restaurar backup",str(Path.home()),"Base SQLite (*.db);;Todos los archivos (*)")
        if not path:return
        if QMessageBox.warning(self,"Restaurar backup","Esto reemplazará los datos actuales por los del backup. ¿Continuar?",QMessageBox.StandardButton.Yes|QMessageBox.StandardButton.No)!=QMessageBox.StandardButton.Yes:return
        try:self.db.restore_backup(path); self.data_changed.emit(); QMessageBox.information(self,"Restaurado","Base restaurada correctamente.")
        except Exception as e: QMessageBox.critical(self,"No se pudo restaurar",str(e))
    def export_excel(self):
        # OpenPyXL es pesado; sólo se carga cuando el usuario realmente exporta.
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        default=f"AppGastos_export_{datetime.now():%Y-%m-%d}.xlsx"; path,_=QFileDialog.getSaveFileName(self,"Exportar datos",str(Path.home()/default),"Excel (*.xlsx)")
        if not path:return
        try:
            wb=Workbook(); ws=wb.active; ws.title="Movimientos"; headers=["Fecha","Tipo","Descripción","Categoría","Cuenta","Cuenta destino","Importe","Origen","Etiquetas","Nota"]
            ws.append(headers)
            for tx in self.db.transactions(): ws.append([tx["tx_date"],tx["kind"],tx["description"],tx["category_display"],tx["account_name"],tx["to_account_name"] or "",float(tx["amount"]),tx.get("source","manual"),tx["tags"],tx["note"]])
            acc=wb.create_sheet("Cuentas"); acc.append(["Cuenta","Tipo","Saldo inicial","Saldo actual","Incluida en saldo disponible"])
            for a in self.db.accounts_with_balances(): acc.append([a["name"],a["type"],float(a["opening_balance"]),float(a["balance"]),"Sí" if a.get("include_in_balance",1) else "No"])
            adj=wb.create_sheet("Ajustes de cuenta"); adj.append(["Fecha","Cuenta","Tipo","Descripción","Importe"])
            for r in self.db.account_adjustments(): adj.append([r["adjustment_date"],r["account_name"],r["adjustment_type"],r["description"],float(r["amount"])])
            cat=wb.create_sheet("Categorías"); cat.append(["Tipo","Ícono","Principal","Subcategoría","Color"])
            for c in self.db.categories(): cat.append([c["kind"],c.get("icon") or c.get("parent_icon") or "other",c["parent_name"] or c["name"],c["name"] if c["parent_name"] else "",c["color"]])
            imp=wb.create_sheet("Importaciones"); imp.append(["Fuente","Fecha","Tipo","Descripción","Importe","Cuenta","Estado","Categoría"])
            for r in self.db.imported_movements("all"): imp.append([r["source"],r["tx_date"],r["kind"],r["description"],float(r["amount"]),r["account_name"],r["status"],r["category_display"]])
            hist=wb.create_sheet("Historial mensual"); hist.append(["Año","Mes","Tipo","Categoría","Subcategoría","Importe","Archivo origen"])
            for r in self.db.historical_monthly_rows(): hist.append([int(r["year"]),int(r["month"]),r["kind"],r["category_name"],r["subcategory_name"],float(r["amount"]),r["file_name"]])
            flex=wb.create_sheet("Flex envíos"); flex.append(["Semana","Zona","Cantidad","Tarifa aplicada","Subtotal","Registrado"])
            for r in self.db.flex_deliveries_rows(): flex.append([r["week_start"],r["zone_name"],int(r["quantity"]),float(r["unit_price"]),float(r["quantity"])*float(r["unit_price"]),r["created_at"]])
            rates=wb.create_sheet("Flex tarifas"); rates.append(["Zona","Tarifa","Vigente desde","Creada"])
            for r in self.db.flex_rates_rows(): rates.append([r["zone_name"],float(r["price"]),r["effective_from"],r["created_at"]])
            for sheet in wb.worksheets:
                for cell in sheet[1]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="6C63FF")
                for col in range(1,sheet.max_column+1):
                    width=min(45,max(12,max(len(str(sheet.cell(r,col).value or "")) for r in range(1,min(sheet.max_row,500)+1))+2)); sheet.column_dimensions[get_column_letter(col)].width=width
                sheet.freeze_panes="A2"
            wb.save(path); QMessageBox.information(self,"Exportación","Excel creado correctamente.")
        except Exception as e: QMessageBox.critical(self,"Exportación",str(e))
