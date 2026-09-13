from __future__ import annotations

import csv
import hashlib
import io
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any


DATE_KEYS = [
    "TRANSACTION_DATE_SHORT", "SETTLEMENT_DATE_SHORT", "MONEY_RELEASE_DATE_SHORT",
    "TRANSACTION_DATE", "SETTLEMENT_DATE", "MONEY_RELEASE_DATE", "DATE", "FECHA",
]
AMOUNT_KEYS = [
    "SETTLEMENT_NET_AMOUNT", "TRANSACTION_NET_AMOUNT", "NET_AMOUNT", "AMOUNT",
    "IMPORTE", "MONTO",
]
DESCRIPTION_KEYS = [
    "DESCRIPTION", "DESCRIPCION", "DESCRIPCIÓN", "OPERATION_TYPE", "TRANSACTION_TYPE",
    "PAYMENT_METHOD_TYPE", "EXTERNAL_REFERENCE",
]
ID_KEYS = [
    "SOURCE_ID", "TRANSACTION_ID", "SETTLEMENT_ID", "PAYMENT_ID", "OPERATION_ID", "EXTERNAL_REFERENCE",
]

NEGATIVE_TYPES = {
    "WITHDRAWAL", "PAYOUT", "CHARGEBACK", "DISPUTE", "PAYMENT", "DEBIT",
    "TRANSFER_OUT", "MONEY_OUT", "EXPENSE",
}
POSITIVE_TYPES = {
    "REFUND", "CASHBACK", "CREDIT", "TRANSFER_IN", "MONEY_IN", "INCOME",
}
SKIP_RECORD_TYPES = {"TOTAL", "SUBTOTAL", "INITIAL_AVAILABLE_BALANCE", "INITIAL_BALANCE"}


def _normalize_header(value: Any) -> str:
    text = str(value or "").strip().upper()
    text = text.replace(" ", "_").replace("-", "_")
    text = re.sub(r"_+", "_", text)
    return text


def _parse_decimal(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("$", "").replace("ARS", "").replace(" ", "")
    if not text:
        return None
    # MP suele exportar decimales con punto; también toleramos formato AR 1.234,56.
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        parts = text.split(",")
        if len(parts[-1]) <= 2:
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def _parse_date(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text:
        return None
    candidates = [text[:10], text]
    for candidate in candidates:
        try:
            return date.fromisoformat(candidate).isoformat()
        except Exception:
            pass
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d/%m/%y"):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat()
        except Exception:
            pass
    # ISO con hora y zona.
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date().isoformat()
    except Exception:
        return None


def _first(row: dict[str, Any], keys: list[str]):
    for key in keys:
        value = row.get(key)
        if value not in (None, ""):
            return value
    return None


def _signed_amount(row: dict[str, Any]) -> float | None:
    # El reporte de dinero en cuenta recomienda SETTLEMENT_NET_AMOUNT como impacto real.
    for key in AMOUNT_KEYS:
        if key in row:
            amount = _parse_decimal(row.get(key))
            if amount is not None and amount != 0:
                return amount

    credit = _parse_decimal(row.get("NET_CREDIT_AMOUNT") or row.get("NET_CREDIT"))
    debit = _parse_decimal(row.get("NET_DEBIT_AMOUNT") or row.get("NET_DEBIT"))
    if credit is not None or debit is not None:
        return float(credit or 0) - abs(float(debit or 0))

    gross = _parse_decimal(row.get("TRANSACTION_AMOUNT"))
    if gross is None or gross == 0:
        return None
    tx_type = str(row.get("TRANSACTION_TYPE") or row.get("RECORD_TYPE") or "").upper()
    if tx_type in NEGATIVE_TYPES:
        return -abs(gross)
    if tx_type in POSITIVE_TYPES:
        return abs(gross)
    return gross


def _description(row: dict[str, Any]) -> str:
    # Priorizamos una descripción humana. Cuando el reporte solo trae códigos
    # técnicos, traducimos el tipo de operación para que la bandeja Por revisar
    # siga siendo entendible.
    for key in ("DESCRIPTION", "DESCRIPCION", "DESCRIPCIÓN"):
        value = str(row.get(key) or "").strip()
        if value:
            return value

    reference = str(row.get("EXTERNAL_REFERENCE") or "").strip()
    if reference and not reference.isdigit():
        return reference

    tx_type = str(row.get("TRANSACTION_TYPE") or row.get("OPERATION_TYPE") or "").strip().upper()
    labels = {
        "SETTLEMENT": "Pago / movimiento Mercado Pago",
        "REFUND": "Devolución Mercado Pago",
        "CASHBACK": "Cashback Mercado Pago",
        "WITHDRAWAL": "Transferencia desde Mercado Pago",
        "WITHDRAWAL_CANCEL": "Transferencia cancelada Mercado Pago",
        "PAYOUT": "Retiro Mercado Pago",
        "CHARGEBACK": "Contracargo Mercado Pago",
        "DISPUTE": "Reclamo Mercado Pago",
        "SETTLEMENT_SHIPPING": "Cargo de envío Mercado Pago",
        "REFUND_SHIPPING": "Devolución de envío Mercado Pago",
    }
    if tx_type in labels:
        return labels[tx_type]

    method = str(row.get("PAYMENT_METHOD_TYPE") or row.get("PAYMENT_METHOD") or "").strip()
    if method:
        return f"Movimiento Mercado Pago · {method.replace('_', ' ')}"
    return "Movimiento Mercado Pago"


def _external_id(row: dict[str, Any], tx_date: str, amount: float, description: str) -> str:
    values = []
    for key in ID_KEYS:
        value = str(row.get(key) or "").strip()
        if value and value not in values:
            values.append(value)
    if values:
        raw = "|".join(values[:3])
        return "mp:" + hashlib.sha1(raw.encode("utf-8", "ignore")).hexdigest()
    canonical = f"{tx_date}|{amount:.2f}|{description}|{sorted((k,str(v)) for k,v in row.items() if v not in (None,''))}"
    return "mp:" + hashlib.sha1(canonical.encode("utf-8", "ignore")).hexdigest()


def _convert_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for raw in rows:
        row = {_normalize_header(k): v for k, v in raw.items() if k is not None}
        record_type = str(row.get("RECORD_TYPE") or "").strip().upper()
        if record_type in SKIP_RECORD_TYPES:
            continue
        amount = _signed_amount(row)
        if amount is None or abs(amount) < 0.000001:
            continue
        tx_date = None
        for key in DATE_KEYS:
            tx_date = _parse_date(row.get(key))
            if tx_date:
                break
        if not tx_date:
            continue
        desc = _description(row)
        kind = "income" if amount > 0 else "expense"
        out.append({
            "external_id": _external_id(row, tx_date, amount, desc),
            "tx_date": tx_date,
            "kind": kind,
            "amount": abs(float(amount)),
            "description": desc,
            "raw": row,
        })
    return out


def _read_csv_bytes(data: bytes) -> list[dict[str, Any]]:
    text = None
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("No pude leer la codificación del CSV.")
    sample = text[:10000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delimiter = dialect.delimiter
    except Exception:
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    return [dict(row) for row in reader]


def _read_csv(path: Path) -> list[dict[str, Any]]:
    return _read_csv_bytes(path.read_bytes())


def _read_xlsx_bytes(data: bytes) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    iterator = ws.iter_rows(values_only=True)
    try:
        headers = next(iterator)
    except StopIteration:
        return []
    keys = [_normalize_header(v) for v in headers]
    rows = []
    for values in iterator:
        rows.append({keys[i]: values[i] if i < len(values) else None for i in range(len(keys))})
    return rows


def _read_xlsx(path: Path) -> list[dict[str, Any]]:
    return _read_xlsx_bytes(path.read_bytes())


def import_mercadopago_bytes(data: bytes, file_name: str = "report.csv") -> list[dict[str, Any]]:
    suffix = Path(file_name).suffix.lower() or ".csv"
    if suffix == ".csv":
        raw = _read_csv_bytes(data)
    elif suffix in {".xlsx", ".xlsm"}:
        raw = _read_xlsx_bytes(data)
    else:
        # La API normalmente devuelve CSV. Si el nombre no trae extensión,
        # intentamos CSV antes de rechazar el contenido.
        try:
            raw = _read_csv_bytes(data)
        except Exception as exc:
            raise ValueError("Formato de reporte de Mercado Pago no soportado.") from exc
    result = _convert_rows(raw)
    if not result:
        raise ValueError(
            "No encontré movimientos compatibles en el reporte recibido de Mercado Pago."
        )
    return result


def import_mercadopago_report(path: str | Path) -> list[dict[str, Any]]:
    """Lee un reporte de Mercado Pago y devuelve movimientos normalizados.

    Soporta CSV/XLSX y varias versiones de encabezados. Priorizamos
    SETTLEMENT_NET_AMOUNT porque es el campo que representa el impacto real
    sobre el dinero de la cuenta en el reporte oficial de Dinero en cuenta.
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)
    return import_mercadopago_bytes(path.read_bytes(), path.name)

# ---------------------------------------------------------------------------
# Historial mensual desde el Excel personal anterior
# ---------------------------------------------------------------------------

def _legacy_month_from_header(value: Any, fallback_month: int) -> tuple[int, int]:
    """Convierte los encabezados D:O del Excel viejo a (año, mes)."""
    if isinstance(value, datetime):
        return value.year, value.month
    if isinstance(value, date):
        return value.year, value.month
    if isinstance(value, (int, float)):
        try:
            from openpyxl.utils.datetime import from_excel
            converted = from_excel(value)
            if isinstance(converted, datetime):
                return converted.year, converted.month
            if isinstance(converted, date):
                return converted.year, converted.month
        except Exception:
            pass
    text = str(value or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%m/%Y", "%Y/%m"):
        try:
            parsed = datetime.strptime(text, fmt)
            return parsed.year, parsed.month
        except Exception:
            pass
    # Este formato se usó como hoja anual. Si un encabezado no se puede leer,
    # mantenemos un fallback razonable para no desplazar columnas.
    return date.today().year, fallback_month


def _legacy_sheet_categories(ws, kind: str) -> list[dict[str, Any]]:
    """Extrae la estructura de categorías del Excel, incluso si una fila no tiene importes."""
    paths: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str]] = set()
    current_group = ""
    stop_names = {"gasto total", "total mensual", "total", "total general"}

    for row in range(3, ws.max_row + 1):
        group = str(ws.cell(row, 1).value or "").strip()
        detail = str(ws.cell(row, 3).value or "").strip()

        if group:
            if _normalize_header(group).replace("_", " ").lower() in stop_names:
                current_group = ""
                continue
            current_group = group
            key = (kind, current_group, "")
            if key not in seen:
                seen.add(key)
                paths.append({"kind": kind, "category_name": current_group, "subcategory_name": ""})
            continue

        if not current_group or not detail:
            continue
        if _normalize_header(detail) in {"TOTAL_AL_MES", "TOTAL", "PROMEDIO"}:
            continue
        key = (kind, current_group, detail)
        if key not in seen:
            seen.add(key)
            paths.append({"kind": kind, "category_name": current_group, "subcategory_name": detail})

    return paths


def _legacy_sheet_records(ws, kind: str) -> list[dict[str, Any]]:
    """Extrae únicamente filas de detalle; nunca los subtotales de grupo."""
    months: list[tuple[int, int]] = []
    for col in range(4, 16):  # D:O
        months.append(_legacy_month_from_header(ws.cell(2, col).value, col - 3))

    records: dict[tuple, float] = {}
    current_group = ""
    stop_names = {"gasto total", "total mensual", "total", "total general"}

    for row in range(3, ws.max_row + 1):
        group = str(ws.cell(row, 1).value or "").strip()
        detail = str(ws.cell(row, 3).value or "").strip()

        if group:
            if _normalize_header(group).replace("_", " ").lower() in stop_names:
                current_group = ""
                continue
            current_group = group
            # Las filas con texto en A y "Total al mes" en C son encabezados
            # de categoría. Sus números son subtotales y duplicarían el detalle.
            continue

        if not current_group or not detail:
            continue
        if _normalize_header(detail) in {"TOTAL_AL_MES", "TOTAL", "PROMEDIO"}:
            continue

        for offset, col in enumerate(range(4, 16)):
            amount = _parse_decimal(ws.cell(row, col).value)
            if amount is None or amount <= 0:
                continue
            year, month = months[offset]
            key = (kind, year, month, current_group, detail)
            records[key] = records.get(key, 0.0) + float(amount)

    return [
        {
            "kind": key[0],
            "year": key[1],
            "month": key[2],
            "category_name": key[3],
            "subcategory_name": key[4],
            "amount": amount,
        }
        for key, amount in sorted(records.items(), key=lambda item: (item[0][1], item[0][2], item[0][0], item[0][3], item[0][4]))
    ]


def import_legacy_finance_workbook(path: str | Path) -> dict[str, Any]:
    """Lee el archivo 'Gestion de gastos.xlsx' usado antes de App Gastos.

    El Excel contiene totales por subcategoría y mes, no operaciones individuales.
    Por eso devuelve *historial mensual* y no inventa fechas de movimientos.
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("El historial anterior debe ser un archivo .xlsx o .xlsm.")

    from openpyxl import load_workbook

    payload = path.read_bytes()
    file_hash = hashlib.sha256(payload).hexdigest()
    wb = load_workbook(io.BytesIO(payload), read_only=False, data_only=True)

    required = {"Gastos", "Ingresos"}
    missing = required.difference(wb.sheetnames)
    if missing:
        raise ValueError("No encontré las hojas Gastos e Ingresos del Excel anterior.")

    category_paths = _legacy_sheet_categories(wb["Gastos"], "expense")
    category_paths.extend(_legacy_sheet_categories(wb["Ingresos"], "income"))

    records = _legacy_sheet_records(wb["Gastos"], "expense")
    records.extend(_legacy_sheet_records(wb["Ingresos"], "income"))
    if not records:
        raise ValueError("No encontré importes mensuales para importar en el archivo.")

    months = sorted({(r["year"], r["month"]) for r in records})
    expense = sum(float(r["amount"]) for r in records if r["kind"] == "expense")
    income = sum(float(r["amount"]) for r in records if r["kind"] == "income")
    categories = {(r["kind"], r["category_name"], r["subcategory_name"]) for r in category_paths}

    return {
        "path": str(path),
        "file_name": path.name,
        "file_hash": file_hash,
        "records": records,
        "category_paths": category_paths,
        "months": months,
        "record_count": len(records),
        "category_count": len(categories),
        "expense_total": expense,
        "income_total": income,
    }
